"""Pruebas de la importación del archivo al catálogo.

Tres propiedades sostienen el proceso y aquí se verifican: idempotencia,
aislamiento de fallos y consentimiento explícito.
"""

from pathlib import Path

import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.storage import LocalStorage
from app.modules.archive.extraction import derive_title, extract_text
from app.modules.archive.importer import ArchiveImporter
from app.modules.archive.models import DocumentStatus
from app.modules.archive.service import ArchiveService
from app.modules.content.models import ConsentScope, ContentItem, ContentType, PublicationStatus


@pytest.fixture
def archive_root(tmp_path: Path) -> Path:
    for theme in ("migrantes", "lgbtq", "mujeres-discapacidad"):
        for folder in ("testimonios", "entrevistas", "videos"):
            (tmp_path / theme / folder).mkdir(parents=True, exist_ok=True)
    return tmp_path


def write(root: Path, relative: str, content: str = "Un testimonio.") -> Path:
    path = root / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return path


async def scan_and_import(session: AsyncSession, root: Path, storage: LocalStorage, **kwargs):
    await ArchiveService(session, root).scan()
    importer = ArchiveImporter(session, root, storage)
    return await importer.import_documents(**kwargs)


# --- Extracción de texto ----------------------------------------------------


def test_extrae_texto_plano(tmp_path: Path) -> None:
    path = tmp_path / "a.txt"
    path.write_text("No dormí en tres días.", encoding="utf-8")
    assert extract_text(path) == "No dormí en tres días."


def test_extrae_word_incluyendo_tablas(tmp_path: Path) -> None:
    """Las transcripciones en Word suelen traer tablas de hablante y texto."""
    from docx import Document

    document = Document()
    document.add_paragraph("Entrevista con la coordinadora")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Coordinadora"
    table.cell(0, 1).text = "Entre trescientas y cuatrocientas."
    path = tmp_path / "entrevista.docx"
    document.save(str(path))

    texto = extract_text(path)

    assert texto is not None
    assert "Entrevista con la coordinadora" in texto
    assert "Coordinadora · Entre trescientas y cuatrocientas." in texto


def test_un_pdf_sin_capa_de_texto_devuelve_none(tmp_path: Path) -> None:
    """Vacío significaría «no dice nada»; `None` significa «aún no sabemos»."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "escaneado.pdf"
    with path.open("wb") as handle:
        writer.write(handle)

    assert extract_text(path) is None


def test_formato_no_extraible_devuelve_none(tmp_path: Path) -> None:
    path = tmp_path / "video.mp4"
    path.write_bytes(b"\x00" * 16)
    assert extract_text(path) is None


@pytest.mark.parametrize(
    ("archivo", "titulo"),
    [
        ("primera-noche.txt", "Primera noche"),
        ("entrevista_albergue_2024.docx", "Entrevista albergue 2024"),
        ("LA-RAMPA.md", "LA RAMPA"),
    ],
)
def test_titulo_legible_desde_el_nombre(archivo: str, titulo: str) -> None:
    """Es un punto de partida editable: el nombre crudo nadie lo corregiría."""
    assert derive_title(Path(archivo)) == titulo


# --- Importación ------------------------------------------------------------


async def test_importa_documentos_detectados(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/testimonios/primera-noche.txt", "No dormí.")
    write(archive_root, "lgbtq/entrevistas/conversacion.txt", "Una conversación.")

    report = await scan_and_import(session, archive_root, storage)

    assert len(report.created) == 2
    assert report.failed == []


async def test_las_piezas_nacen_en_borrador(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/testimonios/a.txt")

    await scan_and_import(session, archive_root, storage)

    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.status is PublicationStatus.DRAFT


async def test_el_consentimiento_por_defecto_es_el_mas_protector(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    """Importar no autoriza a nadie a alimentar al chatbot."""
    write(archive_root, "migrantes/testimonios/a.txt")

    await scan_and_import(session, archive_root, storage)

    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.consent_scope is ConsentScope.NO_AI
    assert item.is_indexable is False


async def test_el_consentimiento_se_declara_explicitamente(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/testimonios/a.txt")

    await scan_and_import(session, archive_root, storage, consent_scope=ConsentScope.PUBLIC)

    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.consent_scope is ConsentScope.PUBLIC


async def test_conserva_la_procedencia(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/testimonios/primera-noche.txt")

    await scan_and_import(session, archive_root, storage)

    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.source_ref == "migrantes/testimonios/primera-noche.txt"
    assert item.imported_at is not None


async def test_el_tipo_viene_de_la_carpeta(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/entrevistas/a.txt")

    await scan_and_import(session, archive_root, storage)

    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.type is ContentType.INTERVIEW


# --- Idempotencia -----------------------------------------------------------


async def test_importar_dos_veces_no_duplica(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/testimonios/a.txt")
    await scan_and_import(session, archive_root, storage)

    segundo = await scan_and_import(session, archive_root, storage)

    total = (await session.execute(select(func.count(ContentItem.id)))).scalar_one()
    assert total == 1
    assert segundo.created == []


async def test_el_documento_queda_marcado_como_importado(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/testimonios/a.txt")

    await scan_and_import(session, archive_root, storage)

    document = (await ArchiveService(session, archive_root).list_documents())[0]
    assert document.status is DocumentStatus.IMPORTED
    assert document.content_item_id is not None


async def test_nombres_repetidos_se_desambiguan(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    """En un archivo real, los nombres repetidos son la norma, no la excepción."""
    write(archive_root, "migrantes/testimonios/la-rampa.txt")
    write(archive_root, "migrantes/entrevistas/la-rampa.md")

    await scan_and_import(session, archive_root, storage)

    slugs = set((await session.execute(select(ContentItem.slug))).scalars().all())
    assert slugs == {"la-rampa", "la-rampa-2"}


# --- Aislamiento de fallos --------------------------------------------------


async def test_un_documento_ilegible_no_interrumpe_el_lote(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    """Con cientos de piezas, un fallo que detiene el lote obliga a reiniciar."""
    (archive_root / "migrantes/testimonios/roto.pdf").write_bytes(b"esto no es un pdf")
    write(archive_root, "migrantes/testimonios/bueno.txt", "Contenido válido.")
    write(archive_root, "lgbtq/testimonios/otro.txt", "Otro contenido.")

    report = await scan_and_import(session, archive_root, storage)

    assert len(report.created) == 2
    assert len(report.failed) == 1
    assert "roto.pdf" in report.failed[0][0]


async def test_el_fallo_registra_su_causa(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    (archive_root / "migrantes/testimonios/roto.pdf").write_bytes(b"no es un pdf")

    await scan_and_import(session, archive_root, storage)

    document = (await ArchiveService(session, archive_root).list_documents())[0]
    assert document.status is DocumentStatus.FAILED
    assert document.error


async def test_un_documento_fallido_se_reintenta(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    """Corregir el archivo y volver a importar debe bastar."""
    path = archive_root / "migrantes/testimonios/roto.pdf"
    path.write_bytes(b"no es un pdf")
    await scan_and_import(session, archive_root, storage)

    path.unlink()
    write(archive_root, "migrantes/testimonios/roto.txt", "Ahora sí legible.")
    report = await scan_and_import(session, archive_root, storage)

    assert len(report.created) == 1


async def test_una_carpeta_de_comunidad_desconocida_falla_con_causa_clara(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "comunidad_inexistente/testimonios/a.txt")

    report = await scan_and_import(session, archive_root, storage)

    assert len(report.failed) == 1
    assert "comunidad" in report.failed[0][1].lower()


# --- Material audiovisual ---------------------------------------------------


async def test_el_audiovisual_se_copia_al_almacenamiento(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    (archive_root / "migrantes/videos/testimonio.mp4").write_bytes(b"\x00" * 64)

    await scan_and_import(session, archive_root, storage)

    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.media_key is not None
    assert await storage.exists(item.media_key)


async def test_el_original_permanece_intacto(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    """Mover el original haría irreversible la importación."""
    path = archive_root / "migrantes/videos/testimonio.mp4"
    path.write_bytes(b"\x00" * 64)

    await scan_and_import(session, archive_root, storage)

    assert path.is_file()
    assert path.read_bytes() == b"\x00" * 64


async def test_el_audiovisual_queda_pendiente_de_transcripcion(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    (archive_root / "migrantes/videos/testimonio.mp4").write_bytes(b"\x00" * 64)

    report = await scan_and_import(session, archive_root, storage)

    assert report.awaiting_transcription == ["migrantes/videos/testimonio.mp4"]
    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.body is None


async def test_el_material_sensible_va_al_ambito_privado(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    from app.modules.content.models import Sensitivity

    (archive_root / "migrantes/videos/testimonio.mp4").write_bytes(b"\x00" * 64)

    await scan_and_import(session, archive_root, storage, sensitivity=Sensitivity.HIGH)

    item = (await session.execute(select(ContentItem))).unique().scalar_one()
    assert item.media_key is not None
    assert item.media_key.startswith("sensitive/")


# --- Alcance del lote -------------------------------------------------------


async def test_se_puede_importar_una_sola_comunidad(
    session: AsyncSession, archive_root: Path, storage: LocalStorage
) -> None:
    write(archive_root, "migrantes/testimonios/a.txt")
    write(archive_root, "lgbtq/testimonios/b.txt")

    report = await scan_and_import(session, archive_root, storage, theme_slug="lgbtq")

    assert report.created == ["lgbtq/testimonios/b.txt"]
