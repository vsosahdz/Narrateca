"""Extracción de texto de los documentos del archivo.

Cubre los formatos en que suelen llegar las transcripciones: texto plano,
Markdown, Word y PDF. Para el material audiovisual no se extrae nada aquí: su
texto proviene de la campaña de transcripción, que es un trabajo aparte.

Cuando la extracción no es posible, se devuelve `None` en lugar de una cadena
vacía. La diferencia importa: vacío significaría "el documento no dice nada",
y lo que ocurre es que todavía no sabemos qué dice.
"""

from pathlib import Path

#: Formatos de los que se extrae texto directamente.
TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
DOCX_SUFFIXES = {".docx"}
PDF_SUFFIXES = {".pdf"}

EXTRACTABLE = TEXT_SUFFIXES | DOCX_SUFFIXES | PDF_SUFFIXES

#: Formatos cuyo texto llega por transcripción, no por extracción.
NEEDS_TRANSCRIPTION = {".mp4", ".mov", ".webm", ".mp3", ".wav", ".m4a"}


class ExtractionError(Exception):
    """El documento tiene un formato extraíble pero no se pudo leer."""


def _from_text(path: Path) -> str:
    # Los archivos del archivo vienen de equipos distintos y no siempre en
    # UTF-8. Se degrada en vez de fallar: perder un acento es preferible a
    # perder el documento entero.
    return path.read_text(encoding="utf-8", errors="replace")


def _from_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    bloques = [parrafo.text.strip() for parrafo in document.paragraphs]

    # Las transcripciones en Word suelen traer tablas con hablante y texto.
    for tabla in document.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells if celda.text.strip()]
            if celdas:
                bloques.append(" · ".join(celdas))

    return "\n\n".join(bloque for bloque in bloques if bloque)


def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    paginas = [(pagina.extract_text() or "").strip() for pagina in reader.pages]
    return "\n\n".join(pagina for pagina in paginas if pagina)


def can_extract(path: Path) -> bool:
    return path.suffix.lower() in EXTRACTABLE


def needs_transcription(path: Path) -> bool:
    return path.suffix.lower() in NEEDS_TRANSCRIPTION


def extract_text(path: Path) -> str | None:
    """Devuelve el texto del documento, o `None` si su formato no lo permite."""
    suffix = path.suffix.lower()

    try:
        if suffix in TEXT_SUFFIXES:
            texto = _from_text(path)
        elif suffix in DOCX_SUFFIXES:
            texto = _from_docx(path)
        elif suffix in PDF_SUFFIXES:
            texto = _from_pdf(path)
        else:
            return None
    except Exception as error:  # noqa: BLE001 - se reporta con la ruta del documento
        raise ExtractionError(f"No se pudo leer «{path.name}»: {error}") from error

    limpio = texto.strip()
    # Un PDF escaneado sin capa de texto devuelve cadena vacía. Eso no es un
    # documento sin contenido: es un documento que necesita reconocimiento
    # óptico, y conviene que el equipo lo vea como tal.
    return limpio or None


def derive_title(path: Path) -> str:
    """Título legible a partir del nombre del archivo.

    Es un punto de partida editable, no una decisión definitiva: el equipo
    corrige el título en el panel. Poner el nombre del archivo tal cual sería
    peor, porque nadie lo corregiría al no parecer un borrador.
    """
    stem = path.stem.replace("_", " ").replace("-", " ")
    stem = " ".join(stem.split())
    return stem[:1].upper() + stem[1:] if stem else path.name
